import re
from pathlib import Path

import openpyxl
import requests


ARQUIVO_ENTRADA = Path("base.xlsx")
ARQUIVO_SAIDA = Path("base_normalizada_viacep2.xlsx")

LIMITE_TESTE = 1800


def consultar_viacep(cep: str) -> dict | None:
    """
    Consulta um CEP no ViaCEP.
    """

    cep_limpo = re.sub(r"\D", "", str(cep))

    if len(cep_limpo) != 8:
        return None

    url = (
        f"https://viacep.com.br/ws/"
        f"{cep_limpo}/json/"
    )

    resposta = requests.get(
        url,
        timeout=15
    )

    resposta.raise_for_status()

    dados = resposta.json()

    if dados.get("erro"):
        return None

    return dados


def extrair_numero(endereco: str) -> str | None:
    """
    Tenta extrair o número do imóvel do endereço original.

    Exemplo:
    SAO PAULO, R PROF MOACIR EIK ALVARO 66 A
    -> 66 A
    """

    if not endereco:
        return None

    endereco = str(endereco).strip()

    resultado = re.search(
        r"\b(\d+[A-Za-z]?(?:\s+[A-Za-z])?)\s*$",
        endereco
    )

    if resultado:
        return resultado.group(1).strip()

    return None


def montar_endereco(
    dados_cep: dict,
    numero: str | None
) -> str:

    logradouro = dados_cep.get(
        "logradouro",
        ""
    ).strip()

    bairro = dados_cep.get(
        "bairro",
        ""
    ).strip()

    cidade = dados_cep.get(
        "localidade",
        ""
    ).strip()

    uf = dados_cep.get(
        "uf",
        ""
    ).strip()

    cep = dados_cep.get(
        "cep",
        ""
    ).strip()

    partes = []

    if logradouro:

        if numero:
            partes.append(
                f"{logradouro}, {numero}"
            )
        else:
            partes.append(logradouro)

    if bairro:
        partes.append(bairro)

    cidade_uf = ""

    if cidade and uf:
        cidade_uf = f"{cidade} - {uf}"
    elif cidade:
        cidade_uf = cidade
    elif uf:
        cidade_uf = uf

    if cidade_uf:
        partes.append(cidade_uf)

    if cep:
        partes.append(cep)

    return ", ".join(partes)


def executar():

    workbook = openpyxl.load_workbook(
        ARQUIVO_ENTRADA
    )

    planilha = workbook.active

    # Colunas atuais:
    #
    # A INSTALACAO
    # B ENDERECO
    # C LATITUDE
    # D LONGITUDE
    # E CEP
    #
    # Novas:
    # F ENDERECO_NORMALIZADO
    # G STATUS_NORMALIZACAO

    planilha.cell(
        row=1,
        column=6,
        value="ENDERECO_NORMALIZADO"
    )

    planilha.cell(
        row=1,
        column=7,
        value="STATUS_NORMALIZACAO"
    )

    processadas = 0

    for linha in range(
        2,
        planilha.max_row + 1
    ):

        instalacao = planilha.cell(
            linha,
            1
        ).value

        endereco_original = planilha.cell(
            linha,
            2
        ).value

        cep = planilha.cell(
            linha,
            5
        ).value

        endereco_existente = planilha.cell(
            linha,
            6
        ).value

        # Já foi normalizado
        if endereco_existente:
            continue

        if processadas >= LIMITE_TESTE:
            break

        processadas += 1

        print(
            f"\n[{processadas}/{LIMITE_TESTE}]"
        )

        print(
            f"Instalação: {instalacao}"
        )

        print(
            f"Original: {endereco_original}"
        )

        print(
            f"CEP: {cep}"
        )

        try:

            dados = consultar_viacep(
                cep
            )

            if not dados:

                planilha.cell(
                    linha,
                    7,
                    "CEP_NAO_ENCONTRADO"
                )

                print(
                    "✗ CEP não encontrado."
                )

                continue

            numero = extrair_numero(
                endereco_original
            )

            print(
                f"Número identificado: "
                f"{numero or 'N/D'}"
            )

            endereco_normalizado = montar_endereco(
                dados,
                numero
            )

            planilha.cell(
                linha,
                6,
                endereco_normalizado
            )

            planilha.cell(
                linha,
                7,
                "OK"
            )

            print(
                f"✓ ViaCEP: "
                f"{dados.get('logradouro')}"
            )

            print(
                f"✓ Normalizado: "
                f"{endereco_normalizado}"
            )

        except Exception as erro:

            planilha.cell(
                linha,
                7,
                f"ERRO: {erro}"
            )

            print(
                f"✗ Erro: {erro}"
            )

        finally:
            # Salva a cada instalação.
            # Se interromper, não perdemos tudo.
            workbook.save(
                ARQUIVO_SAIDA
            )

    workbook.save(
        ARQUIVO_SAIDA
    )

    workbook.close()

    print(
        "\nNormalização finalizada."
    )


if __name__ == "__main__":
    executar()