import time
from pathlib import Path
from urllib.parse import quote_plus

import openpyxl
from playwright.sync_api import sync_playwright


ARQUIVO_ENTRADA = Path("base_google.xlsx")
ARQUIVO_SAIDA = Path("base_normalizada_oficial.xlsx")

LIMITE_TESTE = 1500


from urllib.parse import quote_plus, unquote_plus


def normalizar_endereco(
    page,
    endereco: str
) -> str | None:

    url = (
        "https://www.google.com/maps/search/"
        f"{quote_plus(endereco)}"
    )

    page.goto(
        url,
        wait_until="domcontentloaded",
        timeout=60000
    )

    try:
        page.wait_for_function(
            "() => window.location.href.includes('/maps/place/')",
            timeout=15000
        )
    except Exception:
        pass

    url_atual = page.url

    if "/maps/place/" in url_atual:

        try:
            parte_endereco = (
                url_atual
                .split("/maps/place/", 1)[1]
                .split("/@", 1)[0]
            )

            endereco_normalizado = unquote_plus(
                parte_endereco
            ).strip()

            if endereco_normalizado:
                return endereco_normalizado

        except Exception:
            pass

    print("⚠ Não consegui extrair endereço.")
    print(f"URL atual: {url_atual}")

    return None


def executar():

    workbook = openpyxl.load_workbook(
        ARQUIVO_ENTRADA
    )

    planilha = workbook.active

    # Mantemos as colunas originais:
    #
    # A = INSTALACAO
    # B = ENDERECO
    # C = LATITUDE
    # D = LONGITUDE
    #
    # Novas:
    # E = ENDERECO_NORMALIZADO
    # F = STATUS_NORMALIZACAO

    planilha.cell(
        row=1,
        column=5,
        value="ENDERECO_NORMALIZADO"
    )

    planilha.cell(
        row=1,
        column=6,
        value="STATUS_NORMALIZACAO"
    )

    processadas = 0

    with sync_playwright() as playwright:

        browser = playwright.chromium.launch(
            channel="chrome",
            headless=False
        )

        context = browser.new_context(
            locale="pt-BR"
        )

        page = context.new_page()

        page.goto(
            "https://www.google.com/",
            wait_until="domcontentloaded",
            timeout=60000
        )

        page.wait_for_timeout(3000)

        try:

            for linha in range(
                2,
                planilha.max_row + 1
            ):

                instalacao = planilha.cell(
                    linha,
                    1
                ).value

                endereco_original = planilha.cell(
                    linha,
                    2
                ).value

                endereco_normalizado_existente = planilha.cell(
                    linha,
                    5
                ).value

                # -----------------------------
                # Já foi normalizado
                # -----------------------------

                if endereco_normalizado_existente:
                    print(
                        f"⏭ {instalacao} já normalizado."
                    )
                    continue

                # -----------------------------
                # Limite de teste
                # -----------------------------

                if processadas >= LIMITE_TESTE:
                    break

                if not endereco_original:
                    continue

                processadas += 1

                print(
                    f"\n[{processadas}/{LIMITE_TESTE}] "
                    f"Instalação {instalacao}"
                )

                print(
                    f"Original: {endereco_original}"
                )

                try:

                    endereco_normalizado = normalizar_endereco(
                        page,
                        str(endereco_original)
                    )

                    if endereco_normalizado:

                        planilha.cell(
                            linha,
                            5,
                            endereco_normalizado
                        )

                        planilha.cell(
                            linha,
                            6,
                            "OK"
                        )

                        print(
                            f"✓ Normalizado: "
                            f"{endereco_normalizado}"
                        )

                    else:

                        planilha.cell(
                            linha,
                            6,
                            "NAO_ENCONTRADO"
                        )

                        print(
                            "✗ Endereço normalizado "
                            "não encontrado."
                        )

                except Exception as erro:

                    planilha.cell(
                        linha,
                        6,
                        f"ERRO: {erro}"
                    )

                    print(
                        f"✗ Erro: {erro}"
                    )

                # Salva a cada registro
                workbook.save(
                    ARQUIVO_SAIDA
                )

                # Mesma ideia do seu scraper:
                # pequena pausa entre registros
                time.sleep(1)

        finally:
            context.close()
            browser.close()

    workbook.save(
        ARQUIVO_SAIDA
    )

    workbook.close()

    print(
        "\nNormalização finalizada."
    )


if __name__ == "__main__":
    executar()