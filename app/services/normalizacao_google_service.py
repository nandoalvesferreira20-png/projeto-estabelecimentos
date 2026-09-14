import time
from pathlib import Path
from urllib.parse import (
    quote_plus,
    unquote_plus,
)

import openpyxl

from playwright.sync_api import (
    sync_playwright,
)


# ==========================================
# NORMALIZAR UM ENDEREÇO
# ==========================================

def normalizar_endereco_google(
    page,
    endereco: str,
) -> str | None:

    url = (
        "https://www.google.com/maps/search/"
        f"{quote_plus(endereco)}"
    )

    page.goto(
        url,
        wait_until="domcontentloaded",
        timeout=60000,
    )

    try:

        page.wait_for_function(
            (
                "() => "
                "window.location.href.includes("
                "'/maps/place/'"
                ")"
            ),
            timeout=15000,
        )

    except Exception:
        pass


    url_atual = page.url


    if "/maps/place/" in url_atual:

        try:

            parte_endereco = (
                url_atual
                .split(
                    "/maps/place/",
                    1,
                )[1]
                .split(
                    "/@",
                    1,
                )[0]
            )

            endereco_normalizado = (
                unquote_plus(
                    parte_endereco
                ).strip()
            )

            if endereco_normalizado:

                return endereco_normalizado

        except Exception:
            pass


    print(
        "⚠ Não consegui extrair endereço."
    )

    print(
        f"URL atual: {url_atual}"
    )

    return None


# ==========================================
# EXECUTAR NORMALIZAÇÃO
# ==========================================

def executar_normalizacao_google(
    arquivo_entrada: Path,
    arquivo_saida: Path,
    limite: int | None = None,
) -> dict:

    arquivo_entrada = Path(
        arquivo_entrada
    )

    arquivo_saida = Path(
        arquivo_saida
    )


    # ======================================
    # VALIDAÇÕES
    # ======================================

    if not arquivo_entrada.exists():

        raise FileNotFoundError(
            (
                "Arquivo de entrada "
                "não encontrado: "
                f"{arquivo_entrada}"
            )
        )


    arquivo_saida.parent.mkdir(
        parents=True,
        exist_ok=True,
    )


    # ======================================
    # ABRE EXCEL
    # ======================================

    workbook = (
        openpyxl.load_workbook(
            arquivo_entrada
        )
    )

    planilha = workbook.active


    # ======================================
    # COLUNAS
    #
    # A = INSTALACAO
    # B = ENDERECO
    # C = LATITUDE
    # D = LONGITUDE
    # E = ENDERECO_NORMALIZADO
    # F = STATUS_NORMALIZACAO
    # ======================================

    planilha.cell(
        row=1,
        column=5,
        value="ENDERECO_NORMALIZADO",
    )

    planilha.cell(
        row=1,
        column=6,
        value="STATUS_NORMALIZACAO",
    )


    total_registros = (
        planilha.max_row - 1
    )

    processadas = 0
    normalizadas = 0
    nao_encontradas = 0
    erros = 0
    ignoradas = 0


    # ======================================
    # PLAYWRIGHT
    # ======================================

    with sync_playwright() as playwright:

        browser = (
            playwright.chromium.launch(
                channel="chrome",
                headless=False,
            )
        )

        context = (
            browser.new_context(
                locale="pt-BR"
            )
        )

        page = context.new_page()


        page.goto(
            "https://www.google.com/",
            wait_until="domcontentloaded",
            timeout=60000,
        )

        page.wait_for_timeout(
            3000
        )


        try:

            for linha in range(
                2,
                planilha.max_row + 1,
            ):

                instalacao = (
                    planilha.cell(
                        linha,
                        1,
                    ).value
                )

                endereco_original = (
                    planilha.cell(
                        linha,
                        2,
                    ).value
                )

                existente = (
                    planilha.cell(
                        linha,
                        5,
                    ).value
                )


                # ==========================
                # JÁ NORMALIZADO
                # ==========================

                if existente:

                    ignoradas += 1

                    print(
                        f"⏭ {instalacao} "
                        "já normalizado."
                    )

                    continue


                # ==========================
                # LIMITE
                # ==========================

                if (
                    limite is not None
                    and
                    processadas >= limite
                ):

                    break


                # ==========================
                # SEM ENDEREÇO
                # ==========================

                if not endereco_original:

                    planilha.cell(
                        linha,
                        6,
                        "SEM_ENDERECO",
                    )

                    continue


                processadas += 1


                print(
                    "\n"
                    f"[{processadas}/"
                    f"{limite or total_registros}] "
                    f"Instalação {instalacao}"
                )

                print(
                    "Original: "
                    f"{endereco_original}"
                )


                # ==========================
                # NORMALIZAÇÃO
                # ==========================

                try:

                    endereco_normalizado = (
                        normalizar_endereco_google(
                            page,
                            str(
                                endereco_original
                            ),
                        )
                    )


                    if endereco_normalizado:

                        planilha.cell(
                            linha,
                            5,
                            endereco_normalizado,
                        )

                        planilha.cell(
                            linha,
                            6,
                            "OK",
                        )

                        normalizadas += 1

                        print(
                            "✓ Normalizado: "
                            f"{endereco_normalizado}"
                        )

                    else:

                        planilha.cell(
                            linha,
                            6,
                            "NAO_ENCONTRADO",
                        )

                        nao_encontradas += 1

                        print(
                            "✗ Endereço "
                            "não encontrado."
                        )


                except Exception as erro:

                    erros += 1

                    planilha.cell(
                        linha,
                        6,
                        f"ERRO: {erro}",
                    )

                    print(
                        f"✗ Erro: {erro}"
                    )


                # ==========================
                # SALVA PROGRESSO
                # ==========================

                workbook.save(
                    arquivo_saida
                )


                # ==========================
                # PAUSA
                # ==========================

                time.sleep(
                    1
                )


        finally:

            context.close()
            browser.close()


    # ======================================
    # SALVA RESULTADO FINAL
    # ======================================

    workbook.save(
        arquivo_saida
    )

    workbook.close()


    print(
        "\nNormalização Google finalizada."
    )


    # ======================================
    # RETORNO PARA INTERFACE
    # ======================================

    return {
        "total_registros": total_registros,
        "processadas": processadas,
        "normalizadas": normalizadas,
        "nao_encontradas": nao_encontradas,
        "ignoradas": ignoradas,
        "erros": erros,
        "arquivo_saida": str(
            arquivo_saida
        ),
    }