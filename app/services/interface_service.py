from pathlib import Path

from openpyxl import load_workbook


CAMINHO_RESULTADOS = Path(
    "data/resultados/resultados.xlsx"
)

PASTA_IMAGENS = Path(
    "data/imagens"
)

NOME_PLANILHA = "Análises"


# ==========================================
# BUSCAR IMAGEM DA INSTALAÇÃO
# ==========================================

def buscar_imagem(
    id_instalacao: str,
    data_imagem: str
) -> str | None:

    # --------------------------------------
    # Primeiro tenta ID + data
    #
    # Exemplo:
    # 0056421541_2025-11.*
    # --------------------------------------

    padrao = (
        f"{id_instalacao}_"
        f"{data_imagem}.*"
    )

    arquivos = list(
        PASTA_IMAGENS.glob(
            padrao
        )
    )

    # --------------------------------------
    # Se não encontrar, tenta somente o ID
    # --------------------------------------

    if not arquivos:

        arquivos = list(
            PASTA_IMAGENS.glob(
                f"{id_instalacao}_*"
            )
        )

    # --------------------------------------
    # Nenhuma imagem encontrada
    # --------------------------------------

    if not arquivos:
        return None

    arquivo = arquivos[0]

    return (
        f"/imagens/{arquivo.name}"
    )


# ==========================================
# CARREGAR RESULTADOS
# ==========================================

def carregar_resultados() -> list[dict]:

    if not CAMINHO_RESULTADOS.exists():
        return []

    workbook = load_workbook(
        CAMINHO_RESULTADOS,
        read_only=True,
        data_only=True
    )

    planilha = workbook[
        NOME_PLANILHA
    ]

    resultados = []

    for linha in planilha.iter_rows(
        min_row=2,
        values_only=True
    ):

        (
            id_instalacao,
            data_imagem,
            status,
            confianca,
            motivo_principal,
            evidencias,
            revisao_necessaria,
            observacao_analista,
            data_analise,
        ) = linha

        if not id_instalacao:
            continue

        # ==================================
        # NORMALIZA ID
        # ==================================

        id_instalacao = str(
            id_instalacao
        ).strip()

        # ==================================
        # NORMALIZA DATA
        # ==================================

        if data_imagem:

            data_imagem_texto = str(
                data_imagem
            ).strip()

        else:

            data_imagem_texto = "N/D"

        # ==================================
        # EVIDÊNCIAS
        # ==================================

        evidencias_lista = []

        if evidencias:

            evidencias_lista = [
                item.strip()
                for item
                in str(
                    evidencias
                ).split("|")
                if item.strip()
            ]

        # ==================================
        # LOCALIZA IMAGEM REAL
        # ==================================

        imagem = buscar_imagem(
            id_instalacao=id_instalacao,
            data_imagem=data_imagem_texto
        )

        # ==================================
        # MONTA RESULTADO
        # ==================================

        resultado = {

            "id_instalacao":
                id_instalacao,

            "data_imagem":
                data_imagem_texto,

            "status": (
                str(status)
                if status
                else "INCONCLUSIVO"
            ),

            "confianca": (
                float(confianca)
                if confianca is not None
                else 0
            ),

            "motivo_principal": (
                str(motivo_principal)
                if motivo_principal
                else ""
            ),

            "evidencias":
                evidencias_lista,

            "requer_revisao": (
                str(
                    revisao_necessaria
                )
                .strip()
                .upper()
                == "SIM"
            ),

            "observacao_analista": (
                str(
                    observacao_analista
                )
                if observacao_analista
                else ""
            ),

            "data_analise": (
                data_analise.isoformat()
                if hasattr(
                    data_analise,
                    "isoformat"
                )
                else str(
                    data_analise
                )
                if data_analise
                else None
            ),

            "imagem":
                imagem,
        }

        resultados.append(
            resultado
        )

    workbook.close()

    return resultados