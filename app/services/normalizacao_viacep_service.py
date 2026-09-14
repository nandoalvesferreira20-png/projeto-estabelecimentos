import re
from pathlib import Path

import openpyxl
import requests


# ==========================================
# CONSULTA VIACEP
# ==========================================

def consultar_viacep(
    cep: str
) -> dict | None:

    cep_limpo = re.sub(
        r"\D",
        "",
        str(cep)
    )

    if len(cep_limpo) != 8:
        return None

    url = (
        "https://viacep.com.br/ws/"
        f"{cep_limpo}/json/"
    )

    resposta = requests.get(
        url,
        timeout=15
    )

    resposta.raise_for_status()

    dados = resposta.json()

    if dados.get("erro"):
        return None

    return dados


# ==========================================
# EXTRAI NÚMERO
# ==========================================

def extrair_numero(
    endereco: str
) -> str | None:

    if not endereco:
        return None

    endereco = str(
        endereco
    ).strip()

    resultado = re.search(
        r"\b(\d+[A-Za-z]?(?:\s+[A-Za-z])?)\s*$",
        endereco
    )

    if resultado:

        return (
            resultado
            .group(1)
            .strip()
        )

    return None


# ==========================================
# MONTA ENDEREÇO
# ==========================================

def montar_endereco(
    dados_cep: dict,
    numero: str | None
) -> str:

    logradouro = (
        dados_cep.get(
            "logradouro",
            ""
        ).strip()
    )

    bairro = (
        dados_cep.get(
            "bairro",
            ""
        ).strip()
    )

    cidade = (
        dados_cep.get(
            "localidade",
            ""
        ).strip()
    )

    uf = (
        dados_cep.get(
            "uf",
            ""
        ).strip()
    )

    cep = (
        dados_cep.get(
            "cep",
            ""
        ).strip()
    )

    partes = []


    if logradouro:

        if numero:

            partes.append(
                f"{logradouro}, {numero}"
            )

        else:

            partes.append(
                logradouro
            )


    if bairro:

        partes.append(
            bairro
        )


    cidade_uf = ""


    if cidade and uf:

        cidade_uf = (
            f"{cidade} - {uf}"
        )

    elif cidade:

        cidade_uf = cidade

    elif uf:

        cidade_uf = uf


    if cidade_uf:

        partes.append(
            cidade_uf
        )


    if cep:

        partes.append(
            cep
        )


    return ", ".join(
        partes
    )


# ==========================================
# EXECUTAR NORMALIZAÇÃO VIACEP
# ==========================================

def executar_normalizacao_viacep(
    arquivo_entrada: Path,
    arquivo_saida: Path,
    limite: int | None = None,
) -> dict:

    arquivo_entrada = Path(
        arquivo_entrada
    )

    arquivo_saida = Path(
        arquivo_saida
    )


    # ======================================
    # VALIDA ARQUIVO
    # ======================================

    if not arquivo_entrada.exists():

        raise FileNotFoundError(
            (
                "Arquivo de entrada "
                "não encontrado: "
                f"{arquivo_entrada}"
            )
        )


    arquivo_saida.parent.mkdir(
        parents=True,
        exist_ok=True,
    )


    # ======================================
    # CARREGA EXCEL
    # ======================================

    workbook = (
        openpyxl.load_workbook(
            arquivo_entrada
        )
    )

    planilha = workbook.active


    # ======================================
    # COLUNAS
    #
    # A INSTALACAO
    # B ENDERECO
    # C LATITUDE
    # D LONGITUDE
    # E CEP
    #
    # F ENDERECO_NORMALIZADO
    # G STATUS_NORMALIZACAO
    # ======================================

    planilha.cell(
        row=1,
        column=6,
        value="ENDERECO_NORMALIZADO"
    )

    planilha.cell(
        row=1,
        column=7,
        value="STATUS_NORMALIZACAO"
    )


    total_registros = (
        planilha.max_row - 1
    )

    processadas = 0
    normalizadas = 0
    cep_nao_encontrado = 0
    ignoradas = 0
    erros = 0


    # ======================================
    # PROCESSAMENTO
    # ======================================

    for linha in range(
        2,
        planilha.max_row + 1
    ):

        instalacao = (
            planilha.cell(
                linha,
                1
            ).value
        )

        endereco_original = (
            planilha.cell(
                linha,
                2
            ).value
        )

        cep = (
            planilha.cell(
                linha,
                5
            ).value
        )

        endereco_existente = (
            planilha.cell(
                linha,
                6
            ).value
        )


        # ==============================
        # JÁ NORMALIZADO
        # ==============================

        if endereco_existente:

            ignoradas += 1

            continue


        # ==============================
        # LIMITE
        # ==============================

        if (
            limite is not None
            and processadas >= limite
        ):

            break


        processadas += 1


        print(
            "\n"
            f"[{processadas}/"
            f"{limite or total_registros}]"
        )

        print(
            f"Instalação: {instalacao}"
        )

        print(
            f"Original: {endereco_original}"
        )

        print(
            f"CEP: {cep}"
        )


        # ==============================
        # VIACEP
        # ==============================

        try:

            dados = consultar_viacep(
                cep
            )


            if not dados:

                planilha.cell(
                    linha,
                    7,
                    "CEP_NAO_ENCONTRADO"
                )

                cep_nao_encontrado += 1

                print(
                    "✗ CEP não encontrado."
                )

                continue


            numero = extrair_numero(
                endereco_original
            )


            print(
                "Número identificado: "
                f"{numero or 'N/D'}"
            )


            endereco_normalizado = (
                montar_endereco(
                    dados,
                    numero
                )
            )


            planilha.cell(
                linha,
                6,
                endereco_normalizado
            )

            planilha.cell(
                linha,
                7,
                "OK"
            )


            normalizadas += 1


            print(
                "✓ ViaCEP: "
                f"{dados.get('logradouro')}"
            )

            print(
                "✓ Normalizado: "
                f"{endereco_normalizado}"
            )


        except Exception as erro:

            erros += 1

            planilha.cell(
                linha,
                7,
                f"ERRO: {erro}"
            )

            print(
                f"✗ Erro: {erro}"
            )


        finally:

            # Salva a cada registro
            workbook.save(
                arquivo_saida
            )


    # ======================================
    # SALVA RESULTADO FINAL
    # ======================================

    workbook.save(
        arquivo_saida
    )

    workbook.close()


    print(
        "\nNormalização ViaCEP finalizada."
    )


    # ======================================
    # RETORNO PARA INTERFACE
    # ======================================

    return {
        "total_registros":
            total_registros,

        "processadas":
            processadas,

        "normalizadas":
            normalizadas,

        "cep_nao_encontrado":
            cep_nao_encontrado,

        "ignoradas":
            ignoradas,

        "erros":
            erros,

        "arquivo_saida":
            str(
                arquivo_saida
            ),
    }