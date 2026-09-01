from pathlib import Path
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
    Border,
    Side,
)
from openpyxl.utils import get_column_letter


CAMINHO_RESULTADOS = Path(
    "data/resultados/resultados.xlsx"
)

NOME_PLANILHA = "Análises"


CABECALHOS = [
    "ID",
    "Data da imagem",
    "Status",
    "Confiança",
    "Motivo principal",
    "Evidências",
    "Revisão necessária",
    "Observação do analista",
    "Data/Hora da análise",
]


# ==========================================
# CORES
# ==========================================

COR_CABECALHO = "1F4E78"
COR_TEXTO_CABECALHO = "FFFFFF"

COR_ATIVO = "E2F0D9"
COR_INATIVO = "F4CCCC"
COR_INCONCLUSIVO = "FFF2CC"

COR_REVISAO = "FCE5CD"

COR_BORDA = "D9E1F2"


# ==========================================
# CRIAR PLANILHA
# ==========================================

def criar_planilha():

    CAMINHO_RESULTADOS.parent.mkdir(
        parents=True,
        exist_ok=True
    )

    workbook = Workbook()

    planilha = workbook.active
    planilha.title = NOME_PLANILHA

    planilha.append(CABECALHOS)

    # --------------------------------------
    # Cabeçalho
    # --------------------------------------

    preenchimento_cabecalho = PatternFill(
        fill_type="solid",
        fgColor=COR_CABECALHO
    )

    fonte_cabecalho = Font(
        color=COR_TEXTO_CABECALHO,
        bold=True
    )

    borda = Border(
        bottom=Side(
            style="thin",
            color=COR_BORDA
        )
    )

    for celula in planilha[1]:

        celula.fill = preenchimento_cabecalho
        celula.font = fonte_cabecalho

        celula.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

        celula.border = borda

    # Altura do cabeçalho
    planilha.row_dimensions[1].height = 30

    # Congela cabeçalho
    planilha.freeze_panes = "A2"

    # Filtros
    planilha.auto_filter.ref = (
        f"A1:{get_column_letter(len(CABECALHOS))}1"
    )

    # --------------------------------------
    # Largura das colunas
    # --------------------------------------

    larguras = {
        "A": 18,   # ID
        "B": 16,   # Data imagem
        "C": 24,   # Status
        "D": 14,   # Confiança
        "E": 55,   # Motivo
        "F": 70,   # Evidências
        "G": 22,   # Revisão
        "H": 50,   # Observação analista
        "I": 22,   # Data análise
    }

    for coluna, largura in larguras.items():

        planilha.column_dimensions[
            coluna
        ].width = largura

    workbook.save(
        CAMINHO_RESULTADOS
    )


# ==========================================
# SALVAR RESULTADO
# ==========================================

def salvar_resultado(
    id_instalacao: str,
    data_imagem: str,
    status: str,
    confianca: float,
    motivo_principal: str,
    evidencias: list[str],
    requer_revisao: bool
):

    if not CAMINHO_RESULTADOS.exists():
        criar_planilha()

    workbook = load_workbook(
        CAMINHO_RESULTADOS
    )

    planilha = workbook[NOME_PLANILHA]

    evidencias_texto = " | ".join(
        evidencias
    )

    revisao_texto = (
        "SIM"
        if requer_revisao
        else "NÃO"
    )

    planilha.append([
        str(id_instalacao),
        data_imagem,
        status,
        confianca,
        motivo_principal,
        evidencias_texto,
        revisao_texto,
        "",
        datetime.now()
    ])

    linha = planilha.max_row

    # ======================================
    # FORMATAÇÃO DA LINHA
    # ======================================

    for coluna in range(
        1,
        len(CABECALHOS) + 1
    ):

        celula = planilha.cell(
            row=linha,
            column=coluna
        )

        celula.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )

    # --------------------------------------
    # ID como texto
    # preserva zeros à esquerda
    # --------------------------------------

    planilha.cell(
        row=linha,
        column=1
    ).number_format = "@"

    # --------------------------------------
    # Confiança em %
    # --------------------------------------

    planilha.cell(
        row=linha,
        column=4
    ).number_format = "0%"

    # --------------------------------------
    # Data/hora
    # --------------------------------------

    planilha.cell(
        row=linha,
        column=9
    ).number_format = "dd/mm/yyyy hh:mm"

    # ======================================
    # CORES POR STATUS
    # ======================================

    status_normalizado = (
        str(status)
        .strip()
        .upper()
    )

    if status_normalizado == "ATIVO":

        cor_status = COR_ATIVO

    elif status_normalizado == "POSSIVELMENTE_INATIVO":

        cor_status = COR_INATIVO

    else:

        cor_status = COR_INCONCLUSIVO

    preenchimento_status = PatternFill(
        fill_type="solid",
        fgColor=cor_status
    )

    # Pinta Status
    planilha.cell(
        row=linha,
        column=3
    ).fill = preenchimento_status

    # --------------------------------------
    # Revisão necessária
    # --------------------------------------

    if requer_revisao:

        planilha.cell(
            row=linha,
            column=7
        ).fill = PatternFill(
            fill_type="solid",
            fgColor=COR_REVISAO
        )

        planilha.cell(
            row=linha,
            column=7
        ).font = Font(
            bold=True
        )

    # Altura da linha
    planilha.row_dimensions[
        linha
    ].height = 45

    # Atualiza área do filtro
    planilha.auto_filter.ref = (
        f"A1:"
        f"{get_column_letter(len(CABECALHOS))}"
        f"{planilha.max_row}"
    )

    workbook.save(
        CAMINHO_RESULTADOS
    )

    workbook.close()


# ==========================================
# CARREGAR IDS JÁ PROCESSADOS
# ==========================================

def carregar_arquivos_processados() -> set[str]:

    if not CAMINHO_RESULTADOS.exists():
        return set()

    workbook = load_workbook(
        CAMINHO_RESULTADOS,
        read_only=True,
        data_only=True
    )

    planilha = workbook[NOME_PLANILHA]

    ids_processados = set()

    for linha in planilha.iter_rows(
        min_row=2,
        values_only=True
    ):

        id_instalacao = linha[0]

        if id_instalacao:

            ids_processados.add(
                str(id_instalacao).strip()
            )

    workbook.close()

    return ids_processados